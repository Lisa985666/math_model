from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog


# ===================== 1. 文件路径 =====================
BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "附件" / "附件1.xlsx"
TEMPLATE_FILE = BASE_DIR / "附件" / "附件5" / "result1.xlsx"
OUTPUT_FILE = BASE_DIR / "result1_问题1_计算结果.xlsx"

# ===================== 2. 读取附件1 =====================
def read_data():
    workbook = load_workbook(INPUT_FILE, read_only=True, data_only=True)
    sheet = workbook.active

    time_list = []
    price_list = []
    load_list = []
    pv_list = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        time_list.append(str(row[0]))
        price_list.append(float(row[1]))
        load_list.append(float(row[2]))
        pv_list.append(float(row[3]))

    # 附件中的功率单位是kW，每个时间段为10分钟，因此转换为kWh
    delta_t = 10 / 60
    load_kwh = np.array(load_list) * delta_t
    pv_kwh = np.array(pv_list) * delta_t
    price = np.array(price_list)

    if len(time_list) != 144:
        raise ValueError(f"一天应该有144个时间段，当前读取到{len(time_list)}个")

    return time_list, price, load_kwh, pv_kwh, delta_t


# ===================== 3. 建立并求解问题1模型 =====================
def solve_problem1(price, load, pv, delta_t):
    n = len(price)

    # 变量排列：G、C、D、S、W各占n个位置
    # G：购电量，C：充电量，D：放电量，S：储能量，W：弃光量
    G0 = 0
    C0 = n
    D0 = 2 * n
    S0 = 3 * n
    W0 = 4 * n
    variable_count = 5 * n

    eta = 0.90
    initial_storage = 6000.0
    storage_min = 1200.0
    storage_max = 10800.0
    max_charge_or_discharge = 5000 * delta_t

    # 目标函数：最小化全天购电费用
    objective = np.zeros(variable_count)
    objective[G0:G0 + n] = price

    equality_left = []
    equality_right = []

    # 约束1：能量平衡
    # G + PV + D = Load + C + W
    for t in range(n):
        row = np.zeros(variable_count)
        row[G0 + t] = 1
        row[C0 + t] = -1
        row[D0 + t] = 1
        row[W0 + t] = -1
        equality_left.append(row)
        equality_right.append(load[t] - pv[t])

    # 约束2：储能状态变化
    # S_t = S_(t-1) + 0.9*C_t - D_t/0.9
    for t in range(n):
        row = np.zeros(variable_count)
        row[S0 + t] = 1
        row[C0 + t] = -eta
        row[D0 + t] = 1 / eta

        if t == 0:
            right = initial_storage
        else:
            row[S0 + t - 1] = -1
            right = 0

        equality_left.append(row)
        equality_right.append(right)

    # 约束3：24:00储能量回到6000kWh
    row = np.zeros(variable_count)
    row[S0 + n - 1] = 1
    equality_left.append(row)
    equality_right.append(initial_storage)

    # 各变量上下界
    bounds = (
        [(0, None)] * n  # G：购电量
        + [(0, max_charge_or_discharge)] * n  # C：充电量
        + [(0, max_charge_or_discharge)] * n  # D：放电量
        + [(storage_min, storage_max)] * n  # S：储能量
        + [(0, None)] * n  # W：弃光量
    )

    result = linprog(
        c=objective,
        A_eq=np.array(equality_left),
        b_eq=np.array(equality_right),
        bounds=bounds,
        method="highs",
    )

    if not result.success:
        raise RuntimeError("线性规划求解失败：" + result.message)

    solution = result.x
    G = solution[G0:G0 + n]
    C = solution[C0:C0 + n]
    D = solution[D0:D0 + n]
    S = solution[S0:S0 + n]
    W = solution[W0:W0 + n]

    return G, C, D, S, W


# ===================== 4. 写入题目给出的Excel模板 =====================
def write_result(time_list, G, C, D, S):
    workbook = load_workbook(TEMPLATE_FILE)
    purchase_sheet = workbook["计划购电量"]
    storage_sheet = workbook["充放电量"]

    def round4(value):
        return round(float(value), 4)

    # 写入144个十分钟时段的计划购电量
    for i in range(144):
        purchase_sheet.cell(row=i + 2, column=2).value = round4(G[i])
        purchase_sheet.cell(row=i + 2, column=2).number_format = "0.0000"

    # 6个四小时时段，每段包含24个十分钟时段
    for block in range(6):
        start = block * 24
        end = start + 24
        charge_sum = sum(C[start:end])
        discharge_sum = sum(D[start:end])
        storage_sheet.cell(row=block + 2, column=2).value = round4(charge_sum)
        storage_sheet.cell(row=block + 2, column=3).value = round4(discharge_sum)
        storage_sheet.cell(row=block + 2, column=2).number_format = "0.0000"
        storage_sheet.cell(row=block + 2, column=3).number_format = "0.0000"

    # 0:00和24:00的储能量
    storage_sheet["E2"] = 6000.0000
    storage_sheet["E3"] = round4(S[-1])
    storage_sheet["E2"].number_format = "0.0000"
    storage_sheet["E3"].number_format = "0.0000"

    workbook.save(OUTPUT_FILE)


def main():
    time_list, price, load, pv, delta_t = read_data()
    G, C, D, S, W = solve_problem1(price, load, pv, delta_t)
    write_result(time_list, G, C, D, S)

    # 独立检查
    energy_error = np.max(np.abs(G + pv + D - load - C - W))
    storage_error = np.max(
        np.abs(S[1:] - S[:-1] - 0.9 * C[1:] + D[1:] / 0.9)
    )

    print("问题1求解完成")
    print(f"结果文件：{OUTPUT_FILE}")
    print(f"全天购电量：{G.sum():.4f} kWh")
    print(f"全天购电费用：{np.dot(price, G):.4f} 元")
    print(f"最低储能量：{S.min():.4f} kWh")
    print(f"最高储能量：{S.max():.4f} kWh")
    print(f"24:00储能量：{S[-1]:.4f} kWh")
    print(f"最大能量平衡误差：{energy_error:.12f}")
    print(f"最大储能方程误差：{storage_error:.12f}")


if __name__ == "__main__":
    main()
