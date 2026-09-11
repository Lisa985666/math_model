"""问题二：结合保守预测与低价套利的风险调节随机优化模型。

本程序严格采用统一时间索引：附件1、附件2的第一个数据列对应
2025-01-01 00:10--00:20，最后一个数据列对应次日00:00--00:10，功率数据统一乘以 1/6 转为 kWh。

方法：
1. 负荷场景使用最近同类日及同类日历史分位数；同类日为周日--周四、周五--周六。
2. 光伏场景使用最近7天实际光伏的加权均值和分位数，并用实际非零区间约束日照窗口。
3. 每天0:00建立两日随机日前线性规划。当天普通购电量对所有场景相同，
   场景分别拥有储能充放电和紧急购电变量；只执行第一天的普通购电量。
4. 风险权重由1月份严格滚动验证选择，正式统计为2月1日--12月31日。
5. 白天普通购电量不修改，储能按照实际负荷、实际光伏和SOC实时运行；紧急购电不能给储能充电。

运行：
    python solve_q2_hybrid.py --data-dir abc_extracted/C题/附件
"""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import lil_matrix


DT = 1.0 / 6.0
ETA = 0.90
S_INITIAL = 6000.0
S_MIN = 1200.0
S_MAX = 10800.0
MAX_POWER_KW = 5000.0
MAX_ENERGY_KWH = MAX_POWER_KW * DT
PV_WINDOW_DAYS = 7
PV_THRESHOLD_KW = 1.0
PV_DECAY = 0.8
START_DATE = date(2025, 2, 1)
SCENARIO_RISKS = (0.05, 0.10, 0.15, 0.20, 0.25)


def locate_data_dir(project_root: Path, user_dir: str | None) -> Path:
    candidates: list[Path] = []
    if user_dir:
        candidates.append(Path(user_dir))
    roots = [project_root, *project_root.parents, Path.cwd(), *Path.cwd().parents]
    for root in roots:
        candidates.extend([
            root / "abc_extracted" / "C题" / "附件",
            root / "C题" / "附件",
            root / "附件",
            root / "problem2" / "附件",
        ])
    checked: set[Path] = set()
    for item in candidates:
        item = item.resolve()
        if item in checked:
            continue
        checked.add(item)
        if (item / "附件1.xlsx").exists() and (item / "附件2.xlsx").exists():
            print(f"已找到附件目录：{item}")
            return item
    raise FileNotFoundError("找不到附件目录，请用--data-dir指定包含附件1.xlsx和附件2.xlsx的目录。")


def read_data(data_dir: Path) -> dict:
    """读取数据并统一单位。原始功率的第t列对应第t个10分钟区间。"""
    ws1 = load_workbook(data_dir / "附件1.xlsx", data_only=True).active
    rows1 = list(ws1.iter_rows(min_row=2, values_only=True))
    if len(rows1) != 144:
        raise ValueError(f"附件1应有144个时间点，实际为{len(rows1)}。")
    price = np.array([float(row[1]) for row in rows1], dtype=float)
    q1_load_kwh = np.array([float(row[2]) for row in rows1], dtype=float) * DT
    q1_pv_kwh = np.array([float(row[3]) for row in rows1], dtype=float) * DT

    # 题目时间口径：第1个数据列对应00:10--00:20，最后一列对应次日00:00--00:10。
    # 数组顺序不循环平移；附件1、附件2各列仍按同一列直接对应。
    time_labels = []
    for t in range(144):
        start = ((t + 1) * 10) % 1440
        end = ((t + 2) * 10) % 1440
        sh, sm = divmod(start, 60)
        eh, em = divmod(end, 60)
        end_suffix = "+1" if t == 143 else ""
        time_labels.append(f"{sh}:{sm:02d}-{eh}:{em:02d}{end_suffix}")

    wb2 = load_workbook(data_dir / "附件2.xlsx", data_only=True)
    ws_load = wb2["小区负载"]
    ws_pv = wb2["光伏发电实际功率"]
    dates: list[date] = []
    loads: list[list[float]] = []
    pvs: list[list[float]] = []
    for row_idx in range(2, ws_load.max_row + 1):
        d = ws_load.cell(row_idx, 1).value
        if not isinstance(d, datetime):
            continue
        dates.append(d.date())
        loads.append([float(ws_load.cell(row_idx, col).value or 0.0) for col in range(2, 146)])
        pvs.append([float(ws_pv.cell(row_idx, col).value or 0.0) for col in range(2, 146)])
    if len(dates) != 365:
        raise ValueError(f"附件2应有365天，实际为{len(dates)}天。")

    return {
        "price": price,
        "q1_load_kwh": q1_load_kwh,
        "q1_pv_kwh": q1_pv_kwh,
        "time_labels": time_labels,
        "dates": np.array(dates, dtype=object),
        "loads_kwh": np.asarray(loads, dtype=float) * DT,
        "pvs_kwh": np.asarray(pvs, dtype=float) * DT,
    }


def day_class(current_date: date) -> int:
    return 0 if current_date.weekday() in (6, 0, 1, 2, 3) else 1


def date_at_index(data: dict, idx: int) -> date:
    if idx < len(data["dates"]):
        return data["dates"][idx]
    return data["dates"][-1] + timedelta(days=idx - len(data["dates"]) + 1)


def same_class_history(data: dict, target_idx: int, cutoff_idx: int) -> list[int]:
    target_class = day_class(date_at_index(data, target_idx))
    return [
        idx
        for idx in range(min(cutoff_idx, len(data["dates"])) - 1, -1, -1)
        if day_class(data["dates"][idx]) == target_class
    ]


def daylight_window(profile_kwh: np.ndarray) -> tuple[int, int] | None:
    positive = np.flatnonzero(profile_kwh / DT > PV_THRESHOLD_KW)
    if len(positive) == 0:
        return None
    return int(positive[0]), int(positive[-1])


def pv_history(data: dict, cutoff_idx: int) -> list[int]:
    return list(range(max(0, cutoff_idx - PV_WINDOW_DAYS), cutoff_idx))


def apply_daylight_mask(profile: np.ndarray, histories: list[int], data: dict) -> np.ndarray:
    windows = []
    weights = []
    for idx in histories:
        window = daylight_window(data["pvs_kwh"][idx])
        if window is not None:
            age = histories[-1] - idx
            windows.append(window)
            weights.append(np.exp(-PV_DECAY * age))
    result = np.maximum(profile.copy(), 0.0)
    if not windows:
        return result
    win = np.asarray(windows, dtype=float)
    ww = np.asarray(weights, dtype=float)
    sunrise = int(round(np.average(win[:, 0], weights=ww)))
    sunset = int(round(np.average(win[:, 1], weights=ww)))
    mask = np.zeros(144, dtype=bool)
    mask[max(0, sunrise) : min(143, sunset) + 1] = True
    result[~mask] = 0.0
    return result


def scenario_profiles(data: dict, target_idx: int, cutoff_idx: int) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """返回高风险、基准、低风险三个负荷-光伏场景。"""
    if target_idx >= len(data["dates"]):
        # 12月31日的虚拟下一日采用问题一典型日，避免年末把储能全部耗尽。
        return (
            np.array([data["q1_load_kwh"], data["q1_pv_kwh"]]),
            np.array([data["q1_load_kwh"], data["q1_pv_kwh"]]),
            np.array([data["q1_load_kwh"], data["q1_pv_kwh"]]),
        )

    load_hist_idx = same_class_history(data, target_idx, cutoff_idx)
    if not load_hist_idx:
        load_hist = np.array([data["q1_load_kwh"]])
    else:
        load_hist = data["loads_kwh"][load_hist_idx[:5]]
    load_recent = load_hist[0]
    load_low = np.quantile(load_hist, 0.25, axis=0)
    load_high = np.quantile(load_hist, 0.75, axis=0)

    pv_idx = pv_history(data, cutoff_idx)
    if not pv_idx:
        pv_hist = np.array([data["q1_pv_kwh"]])
        pv_mean = data["q1_pv_kwh"].copy()
        pv_low = pv_mean.copy()
        pv_high = pv_mean.copy()
    else:
        pv_hist = data["pvs_kwh"][pv_idx]
        ages = np.array([cutoff_idx - 1 - idx for idx in pv_idx], dtype=float)
        weights = np.exp(-PV_DECAY * ages)
        weights /= weights.sum()
        pv_mean = np.average(pv_hist, axis=0, weights=weights)
        pv_low = np.quantile(pv_hist, 0.25, axis=0)
        pv_high = np.quantile(pv_hist, 0.75, axis=0)
    pv_mean = apply_daylight_mask(pv_mean, pv_idx, data) if pv_idx else pv_mean
    pv_low = apply_daylight_mask(pv_low, pv_idx, data) if pv_idx else pv_low
    pv_high = apply_daylight_mask(pv_high, pv_idx, data) if pv_idx else pv_high

    high_risk = np.array([load_high, pv_low])
    base_case = np.array([load_recent, pv_mean])
    low_risk = np.array([load_low, pv_high])
    return high_risk, base_case, low_risk


def sparse_matrix(rows: list[dict[int, float]], n_cols: int):
    matrix = lil_matrix((len(rows), n_cols), dtype=float)
    for row_idx, row in enumerate(rows):
        for col_idx, value in row.items():
            matrix[row_idx, col_idx] = value
    return matrix.tocsr()


def solve_stochastic_day_ahead(
    price: np.ndarray,
    scenarios: list[tuple[np.ndarray, np.ndarray]],
    scenario_weights: np.ndarray,
    initial_storage: float,
    planning_emergency_scale: float = 1.0,
    return_full: bool = False,
) -> np.ndarray:
    """两日场景LP，返回所有场景共享的当天普通购电计划。"""
    scenario_count = len(scenarios)
    horizon = scenarios[0][0].shape[0]
    n = 144
    block = horizon * n
    # common grid G + each scenario X,Y,C,D,S,E,W
    # X: 普通电直接供负荷；Y: 光伏直接供负荷。
    # 采用能源分配变量后，紧急购电只能进入负荷平衡式，不能直接进入充电平衡式。
    g0 = 0
    scenario0 = block
    per_scenario = 7 * block
    variable_count = block + scenario_count * per_scenario

    objective = np.zeros(variable_count, dtype=float)
    objective[g0 : g0 + block] = np.tile(price, horizon)
    for s in range(scenario_count):
        e0 = scenario0 + s * per_scenario + 5 * block
        objective[e0 : e0 + block] = (
            planning_emergency_scale
            * scenario_weights[s]
            * 5.0
            * np.tile(price, horizon)
        )

    equality_rows: list[dict[int, float]] = []
    equality_rhs: list[float] = []
    for s, (loads, pvs) in enumerate(scenarios):
        base = scenario0 + s * per_scenario
        x0 = base
        y0 = base + block
        c0 = base + 2 * block
        d0 = base + 3 * block
        s0 = base + 4 * block
        e0 = base + 5 * block
        w0 = base + 6 * block
        for h in range(horizon):
            for t in range(n):
                g = g0 + h * n + t
                x = x0 + h * n + t
                y = y0 + h * n + t
                c = c0 + h * n + t
                d = d0 + h * n + t
                e = e0 + h * n + t
                w = w0 + h * n + t
                # 负荷由普通购电、光伏、储能放电和紧急购电共同满足。
                equality_rows.append({x: 1.0, y: 1.0, d: 1.0, e: 1.0})
                equality_rhs.append(float(loads[h, t]))
                # 普通购电和未直接供负荷的光伏只能去充电或弃电。
                equality_rows.append({g: 1.0, x: -1.0, y: -1.0, c: -1.0, w: -1.0})
                equality_rhs.append(float(-pvs[h, t]))

        for h in range(horizon):
            for t in range(n):
                ss = s0 + h * n + t
                c = c0 + h * n + t
                d = d0 + h * n + t
                row = {ss: 1.0, c: -ETA, d: 1.0 / ETA}
                if h == 0 and t == 0:
                    rhs = initial_storage
                elif t == 0:
                    row[s0 + (h - 1) * n + n - 1] = -1.0
                    rhs = 0.0
                else:
                    row[ss - 1] = -1.0
                    rhs = 0.0
                equality_rows.append(row)
                equality_rhs.append(rhs)

    # 充电量不得由紧急购电提供，只能由普通购电和光伏剩余供给提供。
    inequality_rows: list[dict[int, float]] = []
    inequality_rhs: list[float] = []
    for s, (loads, pvs) in enumerate(scenarios):
        base = scenario0 + s * per_scenario
        x0 = base
        for h in range(horizon):
            for t in range(n):
                g = g0 + h * n + t
                x = x0 + h * n + t
                # X不得超过公共普通购电量G，避免把同一份普通购电重复分配。
                inequality_rows.append({x: 1.0, g: -1.0})
                inequality_rhs.append(0.0)

    bounds = [(0.0, None)] * block
    for s, (loads, pvs) in enumerate(scenarios):
        base = scenario0 + s * per_scenario
        # X: 普通购电供负荷；Y: 光伏供负荷；C/D/S/E/W
        bounds += [(0.0, None)] * block
        bounds += [(0.0, None)] * block
        bounds += [(0.0, MAX_ENERGY_KWH)] * block
        bounds += [(0.0, MAX_ENERGY_KWH)] * block
        bounds += [(S_MIN, S_MAX)] * block
        bounds += [(0.0, None)] * block
        bounds += [(0.0, None)] * block
        # 光伏直接供负荷Y不能超过对应场景的光伏量。
        for h in range(horizon):
            for t in range(n):
                bounds[block + s * per_scenario + block + h * n + t] = (0.0, float(pvs[h, t]))

    result = linprog(
        c=objective,
        A_ub=sparse_matrix(inequality_rows, variable_count),
        b_ub=np.asarray(inequality_rhs),
        A_eq=sparse_matrix(equality_rows, variable_count),
        b_eq=np.asarray(equality_rhs),
        bounds=bounds,
        method="highs",
    )
    if not result.success:
        raise RuntimeError(f"随机日前优化失败：{result.message}")
    if return_full:
        return result.x.copy()
    return result.x[g0 : g0 + n].copy()


def simulate_actual_day(grid_plan: np.ndarray, actual_load: np.ndarray, actual_pv: np.ndarray, initial_storage: float) -> dict:
    charge = np.zeros(144, dtype=float)
    discharge = np.zeros(144, dtype=float)
    storage = np.zeros(144, dtype=float)
    emergency = np.zeros(144, dtype=float)
    curtailment = np.zeros(144, dtype=float)
    current = float(initial_storage)
    for t in range(144):
        net = float(grid_plan[t] + actual_pv[t] - actual_load[t])
        if net >= 0.0:
            charge[t] = min(net, MAX_ENERGY_KWH, max(0.0, (S_MAX - current) / ETA))
            current += ETA * charge[t]
            curtailment[t] = max(0.0, net - charge[t])
        else:
            deficit = -net
            discharge[t] = min(deficit, MAX_ENERGY_KWH, max(0.0, ETA * (current - S_MIN)))
            current -= discharge[t] / ETA
            emergency[t] = max(0.0, deficit - discharge[t])
        current = min(S_MAX, max(S_MIN, current))
        storage[t] = current
    return {
        "charge": charge,
        "discharge": discharge,
        "storage": storage,
        "emergency": emergency,
        "curtailment": curtailment,
        "end_storage": float(current),
    }


def make_day_scenarios(data: dict, day_idx: int, scenario_risk: float):
    targets = [day_idx, day_idx + 1]
    current_cases = [scenario_profiles(data, target, day_idx) for target in targets]
    scenario_weights = np.array([scenario_risk, 1.0 - 2.0 * scenario_risk, scenario_risk], dtype=float)
    scenarios = []
    for s in range(3):
        loads = np.vstack([current_cases[h][s][0] for h in range(2)])
        pvs = np.vstack([current_cases[h][s][1] for h in range(2)])
        scenarios.append((loads, pvs))
    return scenarios, scenario_weights


def run_simulation(
    data: dict,
    scenario_risk: float,
    output_start: date | None = START_DATE,
    planning_emergency_scale: float = 1.0,
    stop_idx: int | None = None,
) -> dict:
    start_output_idx = next(idx for idx, d in enumerate(data["dates"]) if d >= output_start) if output_start else 0
    current_storage = S_INITIAL
    output = {key: [] for key in ["dates", "plan_grid_kwh", "charge_kwh", "discharge_kwh", "storage_kwh", "emergency_kwh", "curtailment_kwh", "initial_storage_kwh", "end_storage_kwh"]}
    total_normal = total_emergency = total_emergency_kwh = 0.0
    last_idx = len(data["dates"]) if stop_idx is None else min(int(stop_idx), len(data["dates"]))
    for day_idx, current_date in enumerate(data["dates"][:last_idx]):
        scenarios, weights = make_day_scenarios(data, day_idx, scenario_risk)
        grid = solve_stochastic_day_ahead(
            data["price"], scenarios, weights, current_storage, planning_emergency_scale
        )
        actual = simulate_actual_day(grid, data["loads_kwh"][day_idx], data["pvs_kwh"][day_idx], current_storage)
        total_normal += float(np.dot(data["price"], grid))
        total_emergency += float(np.dot(5.0 * data["price"], actual["emergency"]))
        total_emergency_kwh += float(actual["emergency"].sum())
        if day_idx >= start_output_idx:
            output["dates"].append(current_date.isoformat())
            output["plan_grid_kwh"].append(grid.tolist())
            output["charge_kwh"].append(actual["charge"].tolist())
            output["discharge_kwh"].append(actual["discharge"].tolist())
            output["storage_kwh"].append(actual["storage"].tolist())
            output["emergency_kwh"].append(actual["emergency"].tolist())
            output["curtailment_kwh"].append(actual["curtailment"].tolist())
            output["initial_storage_kwh"].append(float(current_storage))
            output["end_storage_kwh"].append(float(actual["end_storage"]))
        current_storage = actual["end_storage"]

    normal_cost = sum(float(np.dot(data["price"], row)) for row in output["plan_grid_kwh"])
    emergency_cost = sum(float(np.dot(5.0 * data["price"], row)) for row in output["emergency_kwh"])
    emergency_kwh = sum(float(np.sum(row)) for row in output["emergency_kwh"])
    output["price"] = data["price"].tolist()
    output["time_labels"] = data["time_labels"]
    output["model"] = {
        "time_mapping": "附件1、附件2第一个数据列对应00:10-00:20，最后一列对应次日00:00-00:10；功率乘1/6转为kWh",
        "load_class_0": "周日、周一、周二、周三、周四",
        "load_class_1": "周五、周六",
        "load_forecast": "最近同类日为基准，并用近5个同类日分位数构造场景",
        "pv_forecast": "近7天实际光伏指数加权均值及25/75分位数场景",
        "pv_decay": PV_DECAY,
        "scenario_risk": scenario_risk,
        "planning_emergency_scale": planning_emergency_scale,
        "day_ahead_horizon_days": 2,
        "terminal_value": "12月31日的虚拟下一日使用问题一典型日",
        "real_time_policy": "普通购电固定，实际盈余充电、实际缺口放电，剩余缺口紧急购电",
    }
    output["summary"] = {
        "output_days": len(output["dates"]),
        "normal_purchase_kwh": float(np.sum(output["plan_grid_kwh"])),
        "normal_purchase_cost_yuan": normal_cost,
        "emergency_purchase_kwh": emergency_kwh,
        "emergency_purchase_cost_yuan": emergency_cost,
        "total_cost_yuan": normal_cost + emergency_cost,
        "average_cost_yuan_per_day": (normal_cost + emergency_cost) / len(output["dates"]),
        "emergency_days": int(sum(np.sum(row) > 1e-8 for row in output["emergency_kwh"])),
        "max_emergency_slot_kwh": float(max((max(row) for row in output["emergency_kwh"]), default=0.0)),
        "warmup_full_year_normal_cost_yuan": total_normal,
        "warmup_full_year_emergency_cost_yuan": total_emergency,
        "warmup_full_year_emergency_kwh": total_emergency_kwh,
    }
    return output


def select_risk(data: dict) -> tuple[tuple[float, float], list[dict]]:
    records = []
    planning_scales = (0.05, 0.10, 0.20, 0.40, 0.60, 1.00)
    for risk in SCENARIO_RISKS:
      for scale in planning_scales:
        # 只用1月1日至1月31日的滚动回放选择参数；2月--12月只作为最终检验区间。
        result = run_simulation(data, risk, output_start=None, planning_emergency_scale=scale, stop_idx=31)
        normal = sum(float(np.dot(data["price"], row)) for row in result["plan_grid_kwh"][:31])
        emergency = sum(float(np.dot(5.0 * data["price"], row)) for row in result["emergency_kwh"][:31])
        records.append({"scenario_risk": risk, "planning_emergency_scale": scale, "january_normal_cost_yuan": normal, "january_emergency_cost_yuan": emergency, "january_total_cost_yuan": normal + emergency})
    chosen_record = min(records, key=lambda x: x["january_total_cost_yuan"])
    return (float(chosen_record["scenario_risk"]), float(chosen_record["planning_emergency_scale"])), records


def locate_template(project_root: Path, user_path: str | None) -> Path:
    candidates = []
    if user_path:
        candidates.append(Path(user_path))
    candidates.extend(
        [
            project_root / "abc_extracted" / "C题" / "附件" / "附件5" / "result2.xlsx",
            project_root / "C题" / "附件" / "附件5" / "result2.xlsx",
        ]
    )
    for item in candidates:
        if item.exists():
            return item
    raise FileNotFoundError("找不到result2.xlsx模板，请用--template指定模板路径。")


def write_result_workbook(result: dict, template_path: Path, output_path: Path) -> None:
    """用题目模板写出完整结果表；模型本身仍只依赖numpy/scipy/openpyxl。"""
    wb = load_workbook(template_path)
    plan_ws = wb["计划购电量"]
    storage_ws = wb["充放电量"]
    emergency_ws = wb["紧急购电量"]
    price = np.asarray(result["price"], dtype=float)
    n = 144
    nd = len(result["dates"])

    interval_labels = []
    for t in range(n):
        start = ((t + 1) * 10) % 1440
        end = ((t + 2) * 10) % 1440
        sh, sm = divmod(start, 60)
        eh, em = divmod(end, 60)
        end_suffix = "+1" if t == 143 else ""
        interval_labels.append(f"{sh}:{sm:02d}-{eh}:{em}{end_suffix}")
    for t, label in enumerate(interval_labels, start=2):
        plan_ws.cell(1, t).value = label
    plan_ws.cell(1, 146).value = "全天购电量"
    plan_ws.cell(1, 147).value = "全天购电费"

    for i, iso in enumerate(result["dates"], start=2):
        grid = np.asarray(result["plan_grid_kwh"][i - 2], dtype=float)
        plan_ws.cell(i, 1).value = datetime.fromisoformat(iso)
        for t, value in enumerate(grid, start=2):
            plan_ws.cell(i, t).value = float(value)
        plan_ws.cell(i, 146).value = float(grid.sum())
        plan_ws.cell(i, 147).value = float(np.dot(grid, price))
    for row in plan_ws.iter_rows(min_row=2, max_row=nd + 1, min_col=1, max_col=1):
        row[0].number_format = "yyyy/m/d"
    for row in plan_ws.iter_rows(min_row=2, max_row=nd + 1, min_col=2, max_col=147):
        for cell in row:
            cell.number_format = "0.0000"

    # 清空模板中的示例行，并完整写入输出期内每一天的6个时段。
    # 输出期为2025/2/1--2025/12/31，共334天、2004条明细记录。
    if storage_ws.max_row > 1:
        storage_ws.delete_rows(2, storage_ws.max_row - 1)
    period_names = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    # 实际时钟的4小时段分组。由于第0列对应00:10--00:20，
    # 00:00--04:00包含最后一列(00:00--00:10)和第0--22列。
    storage_blocks = [[143] + list(range(0, 23))]
    storage_blocks += [list(range(23 + 24 * k, 47 + 24 * k)) for k in range(5)]
    for idx in range(nd):
        start_row = 2 + idx * 6
        for k, period in enumerate(period_names):
            r = start_row + k
            storage_ws.cell(r, 1).value = datetime.fromisoformat(result["dates"][idx]) if k == 0 else None
            storage_ws.cell(r, 2).value = period
            slots = storage_blocks[k]
            storage_ws.cell(r, 3).value = float(sum(result["charge_kwh"][idx][t] for t in slots))
            storage_ws.cell(r, 4).value = float(sum(result["discharge_kwh"][idx][t] for t in slots))
            storage_ws.cell(r, 5).value = "0:00" if k == 0 else ("24:00" if k == 1 else None)
            storage_ws.cell(r, 6).value = (
                float(result["initial_storage_kwh"][idx])
                if k == 0
                else (float(result["end_storage_kwh"][idx]) if k == 1 else None)
            )
    for row in storage_ws.iter_rows(min_row=2, max_row=nd * 6 + 1, min_col=1, max_col=1):
        row[0].number_format = "yyyy/m/d"
    for row in storage_ws.iter_rows(min_row=2, max_row=nd * 6 + 1, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = "0.0000"
    for row in storage_ws.iter_rows(min_row=2, max_row=nd * 6 + 1, min_col=6, max_col=6):
        row[0].number_format = "0.0000"

    emergency_rows = []
    for d, iso in enumerate(result["dates"]):
        row = np.asarray(result["emergency_kwh"][d], dtype=float)
        groups = []
        t = 0
        while t < n:
            if row[t] <= 1e-8:
                t += 1
                continue
            start = t
            amount = 0.0
            while t < n and row[t] > 1e-8:
                amount += float(row[t])
                t += 1
            groups.append((f"{interval_labels[start].split('-')[0]}-{interval_labels[t - 1].split('-')[1]}", amount))
        if not groups:
            emergency_rows.append([datetime.fromisoformat(iso), None, 0.0])
        else:
            for j, (label, amount) in enumerate(groups):
                emergency_rows.append([datetime.fromisoformat(iso) if j == 0 else None, label, amount])
    for row in emergency_ws.iter_rows(min_row=2, max_row=max(2000, len(emergency_rows) + 1), min_col=1, max_col=3):
        for cell in row:
            cell.value = None
    for r, values in enumerate(emergency_rows, start=2):
        for c, value in enumerate(values, start=1):
            emergency_ws.cell(r, c).value = value
        emergency_ws.cell(r, 1).number_format = "yyyy/m/d"
        emergency_ws.cell(r, 3).number_format = "0.0000"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", default=None)
    parser.add_argument("--output-json", default="outputs/q2_hybrid/q2_hybrid_solution.json")
    parser.add_argument("--template", default=None)
    parser.add_argument("--output-xlsx", default=None, help="可选：按附件5模板写出result2.xlsx")
    args = parser.parse_args()
    project_root = Path(__file__).resolve().parent
    data = read_data(locate_data_dir(project_root, args.data_dir))
    (risk, scale), validation = select_risk(data)
    result = run_simulation(data, risk, output_start=START_DATE, planning_emergency_scale=scale)
    result["parameter_validation"] = {"candidates": validation, "chosen": {"scenario_risk": risk, "planning_emergency_scale": scale}}
    output_path = Path(args.output_json)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.output_xlsx:
        write_result_workbook(result, locate_template(project_root, args.template), Path(args.output_xlsx))
    print(json.dumps(result["summary"], ensure_ascii=False, indent=2))
    print("chosen_scenario_risk", risk)
    print("chosen_planning_emergency_scale", scale)
    print("output_json", output_path)


if __name__ == "__main__":
    main()
